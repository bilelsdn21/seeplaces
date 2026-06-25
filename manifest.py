"""
SeePlaces — Booked Excursions List (pickup manifest) generator
=============================================================
Reproduces the SeePlaces "Booked excursions list report" the worker reads:
one worksheet per (excursion + date), with the operational pickup columns and
Total / Total seats footer. Built straight from a downloaded supplier report so
the worker no longer has to pull it from SeePlaces separately.
"""

import os
import re

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SUPPLIER_NAME = "Best Time Travel Tunisia"

HEADERS = ["Name and surname", "Location (hotel)", "Meeting point",
           "Room number", "Pickup time", "Remarks", "Seller",
           "Voucher number", "Booking number", "Adult", "Notes"]

_BOLD   = Font(bold=True)
_HDR_FILL = PatternFill("solid", fgColor="E2E8F0")
_thin   = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _s(row, col, default=""):
    if col not in row or pd.isna(row[col]):
        return default
    return str(row[col]).strip()


def _safe_sheet_name(idx: int, name: str) -> str:
    clean = re.sub(r"[:\\/?*\[\]]", " ", name).strip()
    title = f"{idx} {clean}"
    return title[:31]


def _voucher(row) -> str:
    v = _s(row, "VoucherNumber")
    if not v:
        return ""
    try:
        return str(int(float(v))).zfill(7)   # match SeePlaces zero-padding
    except (ValueError, TypeError):
        return v


def _seller(row) -> str:
    email = _s(row, "Expedient email")
    if email:
        return email
    if _s(row, "Payment method").lower() == "online":
        return "Online"
    return ""


def _remarks(row) -> str:
    phone = _s(row, "Participant phone number") or _s(row, "Payer phone number")
    return f"tel. {phone};" if phone else ""


def _adult(row) -> int:
    for col in ("Adult", "Participants count"):
        val = pd.to_numeric(row.get(col), errors="coerce")
        if pd.notna(val):
            return int(val)
    return 0


def _seats(row) -> int:
    val = pd.to_numeric(row.get("Participants count"), errors="coerce")
    if pd.notna(val):
        return int(val)
    return _adult(row)


def generate(report_path: str, out_path: str) -> dict:
    """Build the manifest workbook. Returns a summary dict."""
    df = pd.read_excel(report_path)
    if "Excursion name" not in df.columns:
        raise ValueError("This file is not a SeePlaces report "
                         "(missing 'Excursion name').")

    # Confirmed bookings only — cancelled don't belong on a pickup manifest.
    status = df["Status"].astype(str).str.strip().str.lower()
    df = df[status == "confirmed"].copy()

    df["_exc"]  = df["Excursion name"].astype(str).str.strip()
    df["_date"] = pd.to_datetime(df["ExcursionDate"], errors="coerce")
    df["_datekey"] = df["_date"].dt.strftime("%Y-%m-%d")

    # One sheet per excursion+date, ordered by date then excursion name.
    groups = sorted(
        df.groupby(["_datekey", "_exc"], dropna=False),
        key=lambda kv: (kv[0][0] or "", kv[0][1] or ""))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheet_count = total_bookings = total_pax = 0

    for i, ((datekey, exc), gdf) in enumerate(groups, 1):
        if not exc or exc.lower() == "nan":
            continue
        ws = wb.create_sheet(_safe_sheet_name(i, exc))
        sheet_count += 1

        # ── header block ──────────────────────────────────────────────
        ws["A1"] = "Itineary:";      ws["B1"] = exc
        ws["A2"] = "Supplier name:"; ws["B2"] = SUPPLIER_NAME
        ws["A3"] = "Excursion date"; ws["B3"] = datekey or ""
        for r in (1, 2, 3):
            ws[f"A{r}"].font = _BOLD

        # ── column headers (row 5) ────────────────────────────────────
        hdr_row = 5
        for ci, h in enumerate(HEADERS, 1):
            c = ws.cell(row=hdr_row, column=ci, value=h)
            c.font = _BOLD
            c.fill = _HDR_FILL
            c.border = _BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")

        # ── booking rows ──────────────────────────────────────────────
        row = hdr_row + 1
        sum_adult = sum_seats = 0
        for _, bk in gdf.iterrows():
            name = f"{_s(bk,'Participant first name')} {_s(bk,'Participant last name')}".strip()
            adult = _adult(bk)
            seats = _seats(bk)
            sum_adult += adult
            sum_seats += seats
            vals = [name, _s(bk, "Pickup name"), _s(bk, "Pickup point"),
                    _s(bk, "RoomNumber"), _s(bk, "PickupTime"), _remarks(bk),
                    _seller(bk), _voucher(bk), _s(bk, "Booking number"),
                    adult, _s(bk, "Notes")]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=ci, value=v)
                c.border = _BORDER
                c.alignment = Alignment(vertical="center")
            row += 1
            total_bookings += 1

        total_pax += sum_seats

        # ── footer: Adult / Total / Total seats ───────────────────────
        row += 1
        ws.cell(row=row, column=2, value="Adult").font = _BOLD
        row += 1
        ws.cell(row=row, column=1, value="Total").font = _BOLD
        ws.cell(row=row, column=2, value=sum_adult).font = _BOLD
        row += 1
        ws.cell(row=row, column=1, value="Total seats").font = _BOLD
        ws.cell(row=row, column=2, value=sum_seats).font = _BOLD

        # ── column widths ─────────────────────────────────────────────
        widths = [22, 30, 16, 13, 12, 24, 30, 15, 15, 8, 16]
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    if sheet_count == 0:
        ws = wb.create_sheet("No bookings")
        ws["A1"] = "No confirmed bookings found in this report."

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)

    return {
        "sheets":   sheet_count,
        "bookings": total_bookings,
        "pax":      total_pax,
        "filename": os.path.basename(out_path),
    }
