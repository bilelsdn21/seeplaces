import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ── Colors ────────────────────────────────────────────────────────────────────
ITAKA_RED  = "FFCCCC"
SYS_BLUE   = "CCE5FF"
MISMATCH_Y = "FFF3CD"
MATCH_ODD  = "FFFFFF"
MATCH_EVEN = "F8FAFC"
HEADER_BG  = "1E293B"
WHITE      = "FFFFFF"

GUIDE_COLORS = {"kasha": "4F46E5", "aneta": "0891B2", "joanna": "059669"}
GUIDE_KW     = {
    "joanna": ["johanna", "kisiel"],
    "kasha":  ["katarzyna", "jaszcz"],
    "aneta":  ["aneta", "dymek"],
}

def _fill(c):  return PatternFill("solid", fgColor=c)
def _font(bold=False, color="1E293B", size=9): return Font(bold=bold, color=color, size=size)
def _align(h="left"): return Alignment(horizontal=h, vertical="center", indent=1)


# ── Loaders ───────────────────────────────────────────────────────────────────

def _tnd_price_col(df):
    cols = list(df.columns)
    for i, col in enumerate(cols):
        if "gross price" in str(col).lower() and i + 1 < len(cols):
            if (df[cols[i+1]].dropna().astype(str).str.strip().str.upper() == "TND").any():
                return col
    raise ValueError("Cannot find TND Gross Price column.")


def load_itaka(path: str, guide: str) -> pd.DataFrame:
    df = pd.read_excel(path)
    df = df[df["Status"].astype(str).str.strip().str.lower() == "confirmed"].copy()

    # Filter to this guide's rows using Expedient name
    if "Expedient name" in df.columns and guide in GUIDE_KW:
        kws = GUIDE_KW[guide]
        mask = df["Expedient name"].astype(str).str.lower().apply(lambda n: any(k in n for k in kws))
        df = df[mask].copy()

    df["s_number"]  = df["Booking number"].astype(str).str.strip()
    df["excursion"] = df["Excursion name"].astype(str).str.strip()
    df["date"]      = pd.to_datetime(df["ExcursionDate"], errors="coerce").dt.strftime("%Y-%m-%d")
    df["hotel"]     = df["Pickup name"].astype(str).str.strip()
    df["price"]     = pd.to_numeric(df[_tnd_price_col(df)], errors="coerce").fillna(0)
    return df[["s_number", "excursion", "date", "hotel", "price"]].reset_index(drop=True)


def load_system(path: str) -> pd.DataFrame:
    raw = pd.read_excel(path, header=None)
    header_row = 0
    for i, row in raw.iterrows():
        if row.astype(str).str.strip().eq("Cust. Name").any():
            header_row = i
            break
    df = pd.read_excel(path, header=header_row)
    df.columns = df.columns.astype(str).str.strip()

    df["s_number"]  = df["Cust. Name"].astype(str).str.strip()
    df["excursion"] = df["Sold|Tour Name"].astype(str).str.strip() if "Sold|Tour Name" in df.columns else ""
    df["date"]      = pd.to_datetime(df["Sold|Date"], errors="coerce").dt.strftime("%Y-%m-%d") if "Sold|Date" in df.columns else ""
    df["hotel"]     = df["Hotel Name"].astype(str).str.strip() if "Hotel Name" in df.columns else ""
    df["price"]     = pd.to_numeric(df["Ticket|Amount"], errors="coerce").fillna(0)

    if "Cancel|Adl" in df.columns:
        df = df[pd.to_numeric(df["Cancel|Adl"], errors="coerce").fillna(0) == 0]

    return df[["s_number", "excursion", "date", "hotel", "price"]].reset_index(drop=True)


# ── Matching ──────────────────────────────────────────────────────────────────

def _match(itaka_df, sys_df):
    all_keys = sorted(set(itaka_df["s_number"]) | set(sys_df["s_number"]))
    results  = []
    for key in all_keys:
        i_rows = itaka_df[itaka_df["s_number"] == key].to_dict("records")
        s_rows = sys_df[sys_df["s_number"]     == key].to_dict("records")
        i_used = [False] * len(i_rows)
        s_used = [False] * len(s_rows)

        for ii, ir in enumerate(i_rows):
            for si, sr in enumerate(s_rows):
                if not i_used[ii] and not s_used[si]:
                    if round(float(ir["price"]), 2) == round(float(sr["price"]), 2):
                        results.append({"itaka": ir, "sys": sr, "status": "match", "note": ""})
                        i_used[ii] = s_used[si] = True

        i_left = [r for j, r in enumerate(i_rows) if not i_used[j]]
        s_left = [r for j, r in enumerate(s_rows) if not s_used[j]]
        for k in range(max(len(i_left), len(s_left))):
            ir = i_left[k] if k < len(i_left) else None
            sr = s_left[k] if k < len(s_left) else None
            if ir and sr:
                diff = round(float(sr["price"]) - float(ir["price"]), 2)
                results.append({"itaka": ir, "sys": sr, "status": "mismatch", "note": f"diff {diff:+.2f} TND"})
            elif ir:
                results.append({"itaka": ir, "sys": None, "status": "itaka_only", "note": "Missing in system"})
            else:
                results.append({"itaka": None, "sys": sr, "status": "sys_only", "note": "Not in Itaka"})
    return results


# ── Sheet writer ──────────────────────────────────────────────────────────────

COLS = ["S-Number", "Excursion", "Date", "Hotel", "Price TND"]
WIDTHS = [13, 38, 13, 28, 12]

def _write_sheet(ws, rows, guide, n_itaka, n_sys):
    n     = len(COLS)
    i_s   = 1          # itaka start col
    gap   = i_s + n    # empty gap col
    s_s   = gap + 1    # system start col
    note  = s_s + n    # note col
    color = GUIDE_COLORS.get(guide, "64748B")

    for j, w in enumerate(WIDTHS, i_s):  ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions[get_column_letter(gap)].width = 2
    for j, w in enumerate(WIDTHS, s_s):  ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions[get_column_letter(note)].width = 22

    row = 1

    # Banner
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=note)
    c = ws.cell(row=row, column=1, value=f"  {guide.upper()}  —  Itaka: {n_itaka}  ·  System: {n_sys}")
    c.font = _font(bold=True, color=WHITE, size=11); c.fill = _fill(color); c.alignment = _align()
    ws.row_dimensions[row].height = 26; row += 1

    # Side labels
    ws.merge_cells(start_row=row, start_column=i_s, end_row=row, end_column=i_s + n - 1)
    c = ws.cell(row=row, column=i_s, value="Itaka Report")
    c.font = _font(bold=True, color=WHITE, size=9); c.fill = _fill("DC2626"); c.alignment = _align("center")
    ws.merge_cells(start_row=row, start_column=s_s, end_row=row, end_column=s_s + n - 1)
    c = ws.cell(row=row, column=s_s, value="System Report")
    c.font = _font(bold=True, color=WHITE, size=9); c.fill = _fill("2563EB"); c.alignment = _align("center")
    ws.row_dimensions[row].height = 18; row += 1

    # Headers
    for j, h in enumerate(COLS, i_s):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _font(bold=True, color=WHITE, size=8); c.fill = _fill(HEADER_BG); c.alignment = _align("center")
    for j, h in enumerate(COLS, s_s):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _font(bold=True, color=WHITE, size=8); c.fill = _fill(HEADER_BG); c.alignment = _align("center")
    c = ws.cell(row=row, column=note, value="Note")
    c.font = _font(bold=True, color=WHITE, size=8); c.fill = _fill(HEADER_BG); c.alignment = _align("center")
    ws.row_dimensions[row].height = 16; row += 1

    matched = mismatch = itaka_only = sys_only = 0

    for idx, entry in enumerate(rows):
        status = entry["status"]; ir = entry["itaka"]; sr = entry["sys"]; nt = entry["note"]
        even   = idx % 2 == 0

        if   status == "match":       i_bg = s_bg = MATCH_ODD if even else MATCH_EVEN; matched += 1
        elif status == "mismatch":    i_bg = s_bg = MISMATCH_Y; mismatch += 1
        elif status == "itaka_only":  i_bg = ITAKA_RED;  s_bg = MATCH_ODD if even else MATCH_EVEN; itaka_only += 1
        else:                         i_bg = MATCH_ODD if even else MATCH_EVEN; s_bg = SYS_BLUE; sys_only += 1

        i_vals = [ir["s_number"], ir["excursion"], ir["date"], ir["hotel"], ir["price"]] if ir else ["","","","",""]
        s_vals = [sr["s_number"], sr["excursion"], sr["date"], sr["hotel"], sr["price"]] if sr else ["","","","",""]
        p_i = i_s + n - 1;  p_s = s_s + n - 1

        for j, v in enumerate(i_vals, i_s):
            c = ws.cell(row=row, column=j, value=v); c.fill = _fill(i_bg); c.font = _font(size=9)
            c.alignment = _align("right" if j == p_i else "left")
            if j == p_i and v != "": c.number_format = "#,##0.00"

        for j, v in enumerate(s_vals, s_s):
            c = ws.cell(row=row, column=j, value=v); c.fill = _fill(s_bg); c.font = _font(size=9)
            c.alignment = _align("right" if j == p_s else "left")
            if j == p_s and v != "": c.number_format = "#,##0.00"

        nb = ITAKA_RED if status=="itaka_only" else SYS_BLUE if status=="sys_only" else MISMATCH_Y if status=="mismatch" else (MATCH_ODD if even else MATCH_EVEN)
        c = ws.cell(row=row, column=note, value=nt)
        c.font = _font(size=8, color="64748B"); c.fill = _fill(nb); c.alignment = _align()
        ws.row_dimensions[row].height = 15; row += 1

    return {"matched": matched, "mismatch": mismatch, "itaka_only": itaka_only, "sys_only": sys_only}


def _write_summary(ws, per_guide):
    ws.column_dimensions["A"].width = 22
    for col in "BCDEF": ws.column_dimensions[col].width = 14

    ws.merge_cells("A1:F1")
    c = ws.cell(row=1, column=1, value="Reconciliation Summary")
    c.font = _font(bold=True, color=WHITE, size=13); c.fill = _fill("1E293B"); c.alignment = _align("center")
    ws.row_dimensions[1].height = 30

    for j, h in enumerate(["Guide","Itaka","System","Matched","Price diff","Missing"], 1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = _font(bold=True, color=WHITE, size=9); c.fill = _fill("334155"); c.alignment = _align("center")
    ws.row_dimensions[2].height = 18

    for r, (guide, data) in enumerate(per_guide.items(), 3):
        s = data["stats"]; color = GUIDE_COLORS.get(guide, "64748B")
        vals = [guide.upper(), data["n_itaka"], data["n_sys"],
                s["matched"], s["mismatch"], s["itaka_only"] + s["sys_only"]]
        bgs  = [color, "F8FAFC", "F8FAFC",
                "CCFFCC" if s["matched"] > 0 else "F8FAFC",
                "FFCCCC" if s["mismatch"] > 0 else "F8FAFC",
                "FFCCCC" if (s["itaka_only"]+s["sys_only"]) > 0 else "F8FAFC"]
        fgs  = [WHITE] + ["1E293B"] * 5
        for j, (v, bg, fg) in enumerate(zip(vals, bgs, fgs), 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = _font(bold=(j==1), color=fg, size=10); c.fill = _fill(bg); c.alignment = _align("center" if j > 1 else "left")
        ws.row_dimensions[r].height = 20


# ── Entry point ───────────────────────────────────────────────────────────────

def reconcile(pairs: list, out_path: str) -> dict:
    """
    pairs: list of {guide, itakaPath, sysPath}
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    per_guide = {}

    for pair in pairs:
        guide      = pair["guide"]
        itaka_path = pair["itakaPath"]
        sys_path   = pair["sysPath"]

        print(f"\n{guide.upper()}:")
        itaka_df = load_itaka(itaka_path, guide)
        sys_df   = load_system(sys_path)
        print(f"  Itaka: {len(itaka_df)} bookings  |  System: {len(sys_df)} entries")

        rows  = _match(itaka_df, sys_df)
        ws    = wb.create_sheet(guide.upper())
        stats = _write_sheet(ws, rows, guide, len(itaka_df), len(sys_df))
        per_guide[guide] = {"n_itaka": len(itaka_df), "n_sys": len(sys_df), "stats": stats}
        print(f"  Matched: {stats['matched']}  Mismatch: {stats['mismatch']}  Itaka only: {stats['itaka_only']}  Sys only: {stats['sys_only']}")

    ws_sum = wb.create_sheet("Summary", 0)
    _write_summary(ws_sum, per_guide)
    wb.active = ws_sum
    wb.save(out_path)
    print(f"\nDone -> {out_path}")
    return {g: d["stats"] for g, d in per_guide.items()}
